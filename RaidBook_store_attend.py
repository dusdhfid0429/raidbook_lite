"""
RaidBook 참석/득템 JSON 미러 (보스·전투 워크북 공용)

이행기(3~4단계) 전략: 쓰기 지점(~50곳)은 아직 Excel에 기록하므로,
xlsx의 mtime이 바뀔 때마다 전체를 JSON으로 재가져오는 **미러**를 둔다.
- 재가져오기 비용 = 기존 읽기 1회와 동일 (지금은 매 읽기마다 전체 로드)
- 이후 읽기는 캐시(ms 단위) → 상호작용 3초 타임아웃의 주범 제거
- xlsx가 어떤 경로로 수정되든 mtime으로 감지 → 분기(불일치) 위험 없음

Data/store/attend/<워크북이름>.json 구조:
{
  "schema_version": 1, "kind": "attend_mirror", "updated_at": "...",
  "source_mtime": 1753100000.123,
  "sheets": { "2026-07-22": [ {"msg_id": "...", "event": "...", "event_time": "HH:MM",
                                "uid": "...", "name": "...", ...} ] },
  "loot": { "<msg_id>": {"boss_name": "...", "boss_time": "...", "loot_text": "...",
                          "editor": "...", "ts": "...", "ann_id": "...", "acquirers": "..."} }
}
"""
import os
import threading

from RaidBook_store_core import JsonStore, store_dir


def _format_time(value):
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    if value is None:
        return ""
    return str(value)


class AttendMirror:
    def __init__(self, excel_path: str, schema: dict, system_sheets: set,
                 loot_sheet: str = "획득 아이템"):
        self.excel_path = excel_path
        self.schema = schema
        self.system_sheets = set(system_sheets)
        self.loot_sheet = loot_sheet
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(excel_path)))
        name = os.path.splitext(os.path.basename(excel_path))[0]
        d = store_dir(base_dir, "attend")
        self._store = JsonStore(
            os.path.join(d, f"{name}.json"),
            default_factory=lambda: {"kind": "attend_mirror", "source_mtime": None,
                                     "sheets": {}, "loot": {}},
            kind="attend_mirror",
        )
        self._import_lock = threading.Lock()

    # ── 가져오기 ──────────────────────────────────────────────────────────
    def _import_all(self) -> dict:
        from openpyxl import load_workbook
        from RaidBook_attendance import iter_attend_rows

        sheets, loot = {}, {}
        if os.path.exists(self.excel_path):
            wb = load_workbook(self.excel_path, read_only=True, data_only=True)
            try:
                for sn in wb.sheetnames:
                    if sn == self.loot_sheet:
                        for row in wb[sn].iter_rows(min_row=2, values_only=True):
                            if not row or not row[0]:
                                continue
                            vals = [("" if v is None else v) for v in
                                    (list(row) + [""] * 8)[:8]]
                            loot[str(vals[0])] = {
                                "boss_name": str(vals[1]), "boss_time": _format_time(vals[2]),
                                "loot_text": str(vals[3]), "editor": str(vals[4]),
                                "ts": str(vals[5]), "ann_id": str(vals[6]),
                                "acquirers": str(vals[7]),
                            }
                        continue
                    if sn in self.system_sheets:
                        continue
                    recs = []
                    for rec in iter_attend_rows(wb[sn], self.schema):
                        out = {}
                        for k, v in rec.items():
                            if v is None:
                                out[k] = ""
                            elif isinstance(v, (str, int, float, bool)):
                                out[k] = v
                            else:
                                out[k] = str(v)  # datetime 등 → 문자열 (JSON 직렬화)
                        out["msg_id"] = str(out.get("msg_id", ""))
                        out["uid"] = str(out.get("uid", "")) if out.get("uid") else ""
                        out["event_time"] = _format_time(rec.get("event_time"))
                        recs.append(out)
                    if recs:
                        sheets[sn] = recs
            finally:
                wb.close()
        return {"sheets": sheets, "loot": loot}

    def _source_mtime(self):
        try:
            return os.path.getmtime(self.excel_path)
        except OSError:
            return None

    def ensure_fresh(self) -> dict:
        """xlsx가 바뀌었으면 재가져오기 후 데이터 반환."""
        data = self._store.read()
        mtime = self._source_mtime()
        if data.get("source_mtime") == mtime and (data.get("sheets") or mtime is None):
            return data
        with self._import_lock:
            data = self._store.read()
            mtime = self._source_mtime()
            if data.get("source_mtime") == mtime and (data.get("sheets") or mtime is None):
                return data
            imported = self._import_all()

            def _apply(d):
                d["sheets"] = imported["sheets"]
                d["loot"] = imported["loot"]
                d["source_mtime"] = mtime
            return self._store.mutate(_apply)

    # ── 조회 ─────────────────────────────────────────────────────────────
    def loot_entry(self, message_id) -> dict | None:
        return self.ensure_fresh().get("loot", {}).get(str(message_id))

    def iter_sheets(self, only_sheet: str = None, system_sheets: set = None):
        data = self.ensure_fresh()
        skip = set(system_sheets or self.system_sheets)
        sheets = data.get("sheets", {})
        if only_sheet is not None and only_sheet in sheets:
            yield only_sheet, sheets[only_sheet]
            return
        for sn, recs in sheets.items():
            if sn in skip:
                continue
            yield sn, recs


_mirrors = {}
_mirrors_lock = threading.Lock()


def get_mirror(excel_path: str, schema: dict, system_sheets: set,
               loot_sheet: str = "획득 아이템") -> AttendMirror:
    key = os.path.abspath(excel_path)
    with _mirrors_lock:
        if key not in _mirrors:
            _mirrors[key] = AttendMirror(excel_path, schema, system_sheets, loot_sheet)
        return _mirrors[key]
