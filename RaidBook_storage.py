# -*- coding: utf-8 -*-
import os

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from RaidBook_common import (
    DailyFileHandler,
    do_daily_snapshot,
    do_rolling_backup,
    ensure_workbook_meta,
    get_workbook_meta,
    is_valid_xlsx,
    restore_latest_backup,
    save_workbook_safely,
    save_workbook_safely_sync,
    write_json_atomic,
)


def apply_header(ws, headers, *, fill=None, font=None, border=None):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font
        if border is not None:
            cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20


def auto_column_width(ws, *, padding=4, max_width=None):
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        width = max_length + padding
        if max_width is not None:
            width = min(width, max_width)
        ws.column_dimensions[col_letter].width = width


def apply_border_all(ws, *, border=None, min_row=2, alignment=None):
    alignment = alignment or Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=min_row):
        for cell in row:
            if border is not None:
                cell.border = border
            cell.alignment = alignment


def ensure_dirs(*paths):
    for path in paths:
        os.makedirs(path, exist_ok=True)
