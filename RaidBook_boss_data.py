# -*- coding: utf-8 -*-
import os

from openpyxl import load_workbook

from RaidBook_attendance import BOSS_SCHEMA, iter_attend_rows, resolve_columns


DEFAULT_BOSS_SYSTEM_SHEETS = {
    "요약",
    "보스 출현 통계",
    "획득 아이템",
    "통계",
    "히스토리",
    "보스별 참석 현황",
    "누적 통계",
    "_meta",
}
DEFAULT_LOOT_SHEET = "획득 아이템"
DEFAULT_LOOT_HEADERS = [
    "메시지 ID",
    "보스명",
    "보스 출현 시간",
    "획득 아이템",
    "마지막 수정자",
    "마지막 수정 시각",
    "득템 메시지 ID",
    "획득자",
]


def _write_headers(ws, headers):
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)


def _format_time(value):
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    if value is None:
        return ""
    return str(value)


def _append_loot(wb, result, msg_id):
    loot_items = []
    if "획득 아이템" in wb.sheetnames:
        for row in wb["획득 아이템"].iter_rows(min_row=2, values_only=True):
            if row[0] and str(row[0]) == str(msg_id) and row[3]:
                loot_items.append(row[3])
    result["loot"] = loot_items
    return result


def _has_meaningful_loot(loot_items):
    for item in loot_items or []:
        text = str(item or "").strip()
        if text and text != "없음":
            return True
    return False


def ensure_loot_sheet(wb, *, sheet_name=DEFAULT_LOOT_SHEET, headers=None, apply_header=None):
    """Ensure the loot sheet exists and return it."""
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    if apply_header:
        apply_header(ws, headers or DEFAULT_LOOT_HEADERS)
    else:
        _write_headers(ws, headers or DEFAULT_LOOT_HEADERS)
    return ws


def _mirror(excel_path, *, system_sheets=None, loot_sheet=DEFAULT_LOOT_SHEET):
    """xlsx mtime 기반 JSON 미러 — 실패 시 None 반환하고 호출부는 Excel 직접 읽기로 폴백."""
    try:
        from RaidBook_store_attend import get_mirror
        return get_mirror(excel_path, BOSS_SCHEMA,
                          set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS),
                          loot_sheet)
    except Exception:
        return None


def load_all_loot(excel_path, *, loot_sheet=DEFAULT_LOOT_SHEET):
    """획득 아이템 시트 전체 {메시지ID: {boss_name, boss_time, loot_text, ...}} (JSON 미러 경유)."""
    if not os.path.exists(excel_path):
        return {}
    m = _mirror(excel_path, loot_sheet=loot_sheet)
    if m is None:
        return {}
    try:
        return dict(m.ensure_fresh().get("loot", {}))
    except Exception:
        return {}


def load_loot_record(excel_path, message_id, *, loot_sheet=DEFAULT_LOOT_SHEET):
    """Return (loot_text, editor, announcement_msg_id) for a boss message ID."""
    if not os.path.exists(excel_path):
        return "", "", ""
    m = _mirror(excel_path, loot_sheet=loot_sheet)
    if m is not None:
        try:
            entry = m.loot_entry(message_id)
            if entry is None:
                return "", "", ""
            return entry.get("loot_text", ""), entry.get("editor", ""), entry.get("ann_id", "")
        except Exception:
            pass
    target = str(message_id)
    try:
        wb = load_workbook(excel_path)
    except Exception:
        return "", "", ""
    try:
        if loot_sheet not in wb.sheetnames:
            return "", "", ""
        ws = wb[loot_sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            if str(row[0]) == target:
                loot_text = row[3] if len(row) > 3 and row[3] else ""
                editor = row[4] if len(row) > 4 and row[4] else ""
                announcement_id = str(row[6]) if len(row) > 6 and row[6] else ""
                return loot_text, editor, announcement_id
        return "", "", ""
    finally:
        wb.close()


def load_loot_acquirers(excel_path, message_id, *, loot_sheet=DEFAULT_LOOT_SHEET):
    """Return the acquirers text (8번째 컬럼) for a boss/combat message ID. 없으면 ''."""
    if not os.path.exists(excel_path):
        return ""
    m = _mirror(excel_path, loot_sheet=loot_sheet)
    if m is not None:
        try:
            entry = m.loot_entry(message_id)
            return (entry or {}).get("acquirers", "") or ""
        except Exception:
            pass
    target = str(message_id)
    try:
        wb = load_workbook(excel_path)
    except Exception:
        return ""
    try:
        if loot_sheet not in wb.sheetnames:
            return ""
        ws = wb[loot_sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            if str(row[0]) == target:
                return row[7] if len(row) > 7 and row[7] else ""
        return ""
    finally:
        wb.close()


def upsert_loot_record(
    wb,
    message_id,
    boss_name,
    boss_time,
    loot_text,
    editor,
    *,
    announcement_msg_id=None,
    now_str="",
    acquirers=None,
    loot_sheet=DEFAULT_LOOT_SHEET,
    loot_headers=None,
    apply_header=None,
):
    """Upsert a loot row by message ID and return the current announcement message ID.

    acquirers: None=변경 안 함, 문자열=8번째 컬럼(획득자)에 저장.
    """
    ws = ensure_loot_sheet(
        wb,
        sheet_name=loot_sheet,
        headers=loot_headers or DEFAULT_LOOT_HEADERS,
        apply_header=apply_header,
    )
    # 기존(7컬럼) 시트면 획득자 헤더를 보강
    if not ws.cell(row=1, column=8).value:
        ws.cell(row=1, column=8, value="획득자")
    target = str(message_id)
    existing_ann = ""
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=1)
        if cell.value and str(cell.value) == target:
            existing_ann_cell = ws.cell(row=row_idx, column=7).value
            existing_ann = str(existing_ann_cell) if existing_ann_cell else ""
            ws.cell(row=row_idx, column=2, value=boss_name)
            ws.cell(row=row_idx, column=3, value=boss_time)
            ws.cell(row=row_idx, column=4, value=loot_text)
            ws.cell(row=row_idx, column=5, value=editor)
            ws.cell(row=row_idx, column=6, value=now_str)
            if announcement_msg_id is not None:
                ann_cell = ws.cell(row=row_idx, column=7, value=str(announcement_msg_id))
                ann_cell.number_format = "@"
                existing_ann = str(announcement_msg_id)
            if acquirers is not None:
                ws.cell(row=row_idx, column=8, value=acquirers)
            return existing_ann

    next_row = ws.max_row + 1
    id_cell = ws.cell(row=next_row, column=1, value=target)
    id_cell.number_format = "@"
    ws.cell(row=next_row, column=2, value=boss_name)
    ws.cell(row=next_row, column=3, value=boss_time)
    ws.cell(row=next_row, column=4, value=loot_text)
    ws.cell(row=next_row, column=5, value=editor)
    ws.cell(row=next_row, column=6, value=now_str)
    if announcement_msg_id is not None:
        ann_cell = ws.cell(row=next_row, column=7, value=str(announcement_msg_id))
        ann_cell.number_format = "@"
        existing_ann = str(announcement_msg_id)
    if acquirers is not None:
        ws.cell(row=next_row, column=8, value=acquirers)
    return existing_ann


def ensure_raw_sheet(wb, sheet_name, raw_headers, *, apply_header=None):
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    active = wb.active
    if (
        len(wb.sheetnames) == 1
        and active.max_row == 1
        and active.max_column == 1
        and active["A1"].value is None
    ):
        ws = active
        ws.title = sheet_name
    else:
        ws = wb.create_sheet(sheet_name)
    if apply_header:
        apply_header(ws, raw_headers)
    else:
        _write_headers(ws, raw_headers)
    return ws


def upsert_attendance_records(
    wb,
    sheet_name,
    message_id,
    boss_name,
    boss_time,
    attendees,
    *,
    id_to_name,
    calculate_points,
    raw_headers,
    apply_header=None,
    now_str="",
    prune=False,
):
    """Append missing attendance rows for a boss message and return the target sheet.
    prune=True: 이 메시지의 행 중 attendees에 없는 참석자 행을 지운다 (참석자 목록이 데이터 기준).
    attendees가 비어 있으면 지우지 않는다 — 조회 실패로 빈 목록이 와도 기존 기록 보존."""
    ws = ensure_raw_sheet(wb, sheet_name, raw_headers, apply_header=apply_header)
    cols = resolve_columns(ws, BOSS_SCHEMA)
    target = str(message_id)
    if prune and attendees and cols.get("msg_id") and cols.get("uid"):
        keep = {str(u) for u in attendees}
        drop = []
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=cols["msg_id"]).value or "") != target:
                continue
            uid = ws.cell(row=r, column=cols["uid"]).value
            if uid in (None, "") or str(uid) not in keep:
                drop.append(r)
        for r in reversed(drop):
            ws.delete_rows(r)
    existing_keys = set()
    existing_message = False
    for rec in iter_attend_rows(ws, BOSS_SCHEMA):
        if str(rec["msg_id"]) != target:
            continue
        existing_message = True
        existing_keys.add(str(rec["uid"]) if rec["uid"] is not None else "")

    def _write_row(values_by_field):
        next_row = ws.max_row + 1
        for field, value in values_by_field.items():
            col = cols.get(field)
            if not col:
                continue
            cell = ws.cell(row=next_row, column=col, value=value)
            if field == "msg_id":
                cell.number_format = "@"

    if not attendees:
        if not existing_message:
            _write_row({"msg_id": target, "event": boss_name, "event_time": boss_time,
                        "uid": "", "joined_at": now_str})
        return ws

    for uid in attendees:
        uid = str(uid)
        if uid in existing_keys:
            continue
        _write_row({"msg_id": target, "event": boss_name, "event_time": boss_time,
                    "uid": uid, "joined_at": now_str,
                    "name": id_to_name.get(uid, uid),
                    "points": calculate_points(boss_name, boss_time)})
        existing_keys.add(uid)
    return ws


def load_attendees_by_message_id(excel_path, message_id, *, system_sheets=None):
    if not os.path.exists(excel_path):
        return []
    m = _mirror(excel_path, system_sheets=system_sheets)
    if m is not None:
        try:
            target = str(message_id)
            attendees = []
            for _sn, recs in m.iter_sheets(system_sheets=set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS)):
                for rec in recs:
                    attendee = rec.get("uid")
                    if rec.get("msg_id") == target and attendee and attendee not in attendees:
                        attendees.append(attendee)
            return attendees
        except Exception:
            pass
    system_sheets = set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS)
    wb = load_workbook(excel_path)
    attendees = []
    target = str(message_id)
    try:
        for sheet_name in wb.sheetnames:
            if sheet_name in system_sheets:
                continue
            for rec in iter_attend_rows(wb[sheet_name], BOSS_SCHEMA):
                attendee = rec.get("uid")
                if str(rec["msg_id"]) == target and attendee and attendee not in attendees:
                    attendees.append(attendee)
        return attendees
    finally:
        wb.close()


def find_boss_for_settlement(
    excel_path,
    sheet_name,
    target_time,
    *,
    system_sheets=None,
    return_all=False,
    exact_time=False,
    same_hour_only=False,
    require_loot=False,
    start_minutes=None,
    end_minutes=None,
):
    """
    보스 정산용 기록 조회.
    return_all=False: 가장 가까운 시간 1개 반환 (기존 동작)
    return_all=True: 해당 시트의 기록을 list로 반환 (settle 선택용)
    exact_time=True: target_time과 HH:MM이 정확히 같은 기록만 반환
    same_hour_only=True: target_time과 같은 HH 시간대 기록만 반환
    require_loot=True: 획득 아이템이 등록된 기록만 반환
    start_minutes/end_minutes: 하루 안의 분 단위 범위(포함)로 필터링
    """
    if not os.path.exists(excel_path):
        return None if not return_all else []
    m = _mirror(excel_path, system_sheets=system_sheets)
    if m is not None:
        try:
            return _find_boss_for_settlement_mirror(
                m, sheet_name, target_time,
                system_sheets=set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS),
                return_all=return_all, exact_time=exact_time,
                same_hour_only=same_hour_only, require_loot=require_loot,
                start_minutes=start_minutes, end_minutes=end_minutes)
        except Exception:
            pass
    system_sheets = set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS)
    wb = load_workbook(excel_path)
    try:
        try:
            th, tm = map(int, target_time.split(":"))
            target_minutes = th * 60 + tm
        except Exception:
            target_minutes = None

        sheets = [sheet_name] if sheet_name in wb.sheetnames else [
            s for s in wb.sheetnames if s not in system_sheets
        ]
        candidates = []
        for sn in sheets:
            if sn in system_sheets:
                continue
            ws = wb[sn]
            rows_by_id = {}
            for rec in iter_attend_rows(ws, BOSS_SCHEMA):
                if not rec.get("event"):
                    continue
                mid = str(rec["msg_id"])
                if mid not in rows_by_id:
                    rows_by_id[mid] = {
                        "boss_name": rec["event"] or "",
                        "boss_time": _format_time(rec.get("event_time")),
                        "attendees": [],
                        "sheet_name": sn,
                    }
                if rec.get("uid"):
                    rows_by_id[mid]["attendees"].append(rec["uid"])
            for mid, data in rows_by_id.items():
                diff = float("inf")
                record_minutes = None
                if target_minutes is not None and data["boss_time"]:
                    try:
                        bh, bm = map(int, data["boss_time"].split(":"))
                        record_minutes = bh * 60 + bm
                        diff = abs(record_minutes - target_minutes)
                    except Exception:
                        pass
                if target_minutes is not None:
                    if exact_time and record_minutes != target_minutes:
                        continue
                    if same_hour_only and (record_minutes is None or record_minutes // 60 != th):
                        continue
                if start_minutes is not None and (record_minutes is None or record_minutes < start_minutes):
                    continue
                if end_minutes is not None and (record_minutes is None or record_minutes > end_minutes):
                    continue
                candidates.append({
                    "msg_id": mid,
                    "diff": diff,
                    **data
                })

        if not candidates:
            return None if not return_all else []

        if return_all:
            # diff 기준 정렬 후 모두 반환
            candidates.sort(key=lambda x: x["diff"])
            results = []
            for c in candidates:
                result = _append_loot(wb, c, c["msg_id"])
                if require_loot and not _has_meaningful_loot(result.get("loot")):
                    continue
                results.append(result)
            return results
        else:
            # 가장 가까운 것 1개만 반환 (기존 동작)
            candidates.sort(key=lambda x: x["diff"])
            for candidate in candidates:
                result = _append_loot(wb, candidate, candidate["msg_id"])
                if require_loot and not _has_meaningful_loot(result.get("loot")):
                    continue
                return result
            return None
    finally:
        wb.close()


def _mirror_loot_list(m, mid):
    entry = m.loot_entry(mid)
    text = (entry or {}).get("loot_text", "")
    return [text] if text else []


def _find_boss_for_settlement_mirror(m, sheet_name, target_time, *, system_sheets,
                                     return_all, exact_time, same_hour_only,
                                     require_loot, start_minutes, end_minutes):
    try:
        th, tm = map(int, target_time.split(":"))
        target_minutes = th * 60 + tm
    except Exception:
        target_minutes = None
        th = None

    candidates = []
    for sn, recs in m.iter_sheets(only_sheet=sheet_name, system_sheets=system_sheets):
        rows_by_id = {}
        for rec in recs:
            if not rec.get("event"):
                continue
            mid = rec["msg_id"]
            if mid not in rows_by_id:
                rows_by_id[mid] = {
                    "boss_name": rec["event"] or "",
                    "boss_time": rec.get("event_time") or "",
                    "attendees": [],
                    "sheet_name": sn,
                }
            if rec.get("uid"):
                rows_by_id[mid]["attendees"].append(rec["uid"])
        for mid, data in rows_by_id.items():
            diff = float("inf")
            record_minutes = None
            if target_minutes is not None and data["boss_time"]:
                try:
                    bh, bm = map(int, data["boss_time"].split(":"))
                    record_minutes = bh * 60 + bm
                    diff = abs(record_minutes - target_minutes)
                except Exception:
                    pass
            if target_minutes is not None:
                if exact_time and record_minutes != target_minutes:
                    continue
                if same_hour_only and (record_minutes is None or record_minutes // 60 != th):
                    continue
            if start_minutes is not None and (record_minutes is None or record_minutes < start_minutes):
                continue
            if end_minutes is not None and (record_minutes is None or record_minutes > end_minutes):
                continue
            candidates.append({"msg_id": mid, "diff": diff, **data})

    if not candidates:
        return None if not return_all else []
    candidates.sort(key=lambda x: x["diff"])
    if return_all:
        results = []
        for c in candidates:
            c["loot"] = _mirror_loot_list(m, c["msg_id"])
            if require_loot and not _has_meaningful_loot(c.get("loot")):
                continue
            results.append(c)
        return results
    for c in candidates:
        c["loot"] = _mirror_loot_list(m, c["msg_id"])
        if require_loot and not _has_meaningful_loot(c.get("loot")):
            continue
        return c
    return None


def find_boss_by_message_id(excel_path, msg_id, *, system_sheets=None):
    if not os.path.exists(excel_path):
        return None
    m = _mirror(excel_path, system_sheets=system_sheets)
    if m is not None:
        try:
            target = str(msg_id)
            skip = set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS)
            for sn, recs in m.iter_sheets(system_sheets=skip):
                rows_by_id = {}
                for rec in recs:
                    if not rec.get("event"):
                        continue
                    mid = rec["msg_id"]
                    if mid not in rows_by_id:
                        rows_by_id[mid] = {
                            "boss_name": rec["event"] or "",
                            "boss_time": rec.get("event_time") or "",
                            "attendees": [],
                            "sheet_name": sn,
                        }
                    if rec.get("uid"):
                        rows_by_id[mid]["attendees"].append(rec["uid"])
                if target in rows_by_id:
                    result = {"msg_id": target, **rows_by_id[target]}
                    result["loot"] = _mirror_loot_list(m, target)
                    return result
            return None
        except Exception:
            pass
    target = str(msg_id)
    system_sheets = set(system_sheets or DEFAULT_BOSS_SYSTEM_SHEETS)
    wb = load_workbook(excel_path)
    try:
        result = None
        for sn in wb.sheetnames:
            if sn in system_sheets:
                continue
            ws = wb[sn]
            rows_by_id = {}
            for rec in iter_attend_rows(ws, BOSS_SCHEMA):
                if not rec.get("event"):
                    continue
                mid = str(rec["msg_id"])
                if mid not in rows_by_id:
                    rows_by_id[mid] = {
                        "boss_name": rec["event"] or "",
                        "boss_time": _format_time(rec.get("event_time")),
                        "attendees": [],
                        "sheet_name": sn,
                    }
                if rec.get("uid"):
                    rows_by_id[mid]["attendees"].append(rec["uid"])
            if target in rows_by_id:
                result = {"msg_id": target, **rows_by_id[target]}
                break
        if result is None:
            return None
        return _append_loot(wb, result, target)
    finally:
        wb.close()
